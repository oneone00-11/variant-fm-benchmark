#!/usr/bin/env python3
"""Build the Supplementary Materials (Tables S1-S5) for the manuscript.

Every number is read directly from results/tables/ (or, for S1, the dataset
record). Nothing is recomputed here. Tables map to the main text as:
  S1 -> Methods 2.2 ("Table S1") + Data availability
  S2 -> Methods 2.3 ("coverage in Table 1 and Supplementary")
  S3 -> Discussion/Limitations ("a priori power calculations (Supplementary)")
  S4 -> Table 2 / Fig 2 (per-gene values behind the meta-analysis)
  S5 -> Results 3.4 (ClinVar AUROC/AUPRC, circularity sensitivity)

Outputs: <outdir>/Table_S1..S5.tsv  and  <outdir>/Supplementary_Materials.docx
Usage:   python supplementary/build_supplementary.py [outdir]
"""
import csv
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
T = ROOT / "results" / "tables"
OUTDIR = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "supplementary"
OUTDIR.mkdir(parents=True, exist_ok=True)

# Display names + the splice-ranking order used in main Table 2.
NAME = {
    "pangolin_score": "Pangolin", "spliceai_ds": "SpliceAI",
    "alphagenome_splice": "AlphaGenome", "cadd_phred": "CADD",
    "gpn_msa_score": "GPN-MSA", "phylop100way": "phyloP",
    "phastcons100way": "phastCons", "nucleotide_transformer": "Nucleotide Transformer",
    "gnomad_af_global": "gnomAD AF (global)", "gnomad_af_popmax": "gnomAD AF (popmax)",
    "alphamissense": "AlphaMissense",
}
ORDER = ["pangolin_score", "spliceai_ds", "alphagenome_splice", "cadd_phred",
         "gpn_msa_score", "phylop100way", "phastcons100way",
         "nucleotide_transformer", "gnomad_af_global", "gnomad_af_popmax",
         "alphamissense"]
GENES = ["BRCA1", "BRCA2", "BARD1", "PALB2", "RAD51C", "VHL", "BAP1"]


def w(name, header, rows):
    """Write a TSV and return (header, rows) for the docx builder."""
    with open(OUTDIR / name, "w", newline="") as f:
        wr = csv.writer(f, delimiter="\t")
        wr.writerow(header)
        wr.writerows(rows)
    return header, rows


# ---- Table S1: functional gold-standard datasets (project record) ----------
# URNs/transcripts/licences match Methods 2.2 exactly.
s1_header = ["gene", "mavedb_urn", "assay_type", "reference_transcript",
             "mapped_via_mane_refseq", "assembly", "reference_pmid_source", "license"]
# reference_pmid_source: primary publications, matched to the manuscript reference
# list. BARD1 and PALB2 have no peer-reviewed primary publication (BARD1: medRxiv
# preprint only), so the MaveDB URN is the authoritative source for those two.
s1_rows = [
    ["BRCA1", "urn:mavedb:00000097-0-2", "SGE", "NM_007294.3", "(native RefSeq)", "GRCh38", "Findlay et al. 2018, Nature 562:217–222", "CC0"],
    ["BRCA2", "urn:mavedb:00001225-a-1", "SGE", "ENST00000380152.8", "NM_000059.4", "GRCh38", "Huang et al. 2025, Nature 638:528–537", "CC0"],
    ["BARD1", "urn:mavedb:00001250-a-2", "SGE", "NM_000465.4", "(native RefSeq)", "GRCh38", "MaveDB score set (see URN); no peer-reviewed primary publication (medRxiv preprint only)", "CC0"],
    ["PALB2", "urn:mavedb:00001259-a-2", "SGE", "NM_024675.4", "(native RefSeq)", "GRCh38", "MaveDB score set (see URN); no peer-reviewed primary publication", "CC0"],
    ["RAD51C", "urn:mavedb:00000673-0-1", "SGE", "ENST00000337432.9", "NM_058216.3", "GRCh38", "Olvera-León et al. 2024, Cell 187:5719–5734", "CC BY 4.0"],
    ["VHL", "urn:mavedb:00000675-a-1", "SGE", "ENST00000256474.3", "NM_000551.4", "GRCh38", "Buckley et al. 2024, Nat. Genet. 56:1446–1455", "CC BY 4.0"],
    ["BAP1", "urn:mavedb:00000662-0-1", "SGE", "ENST00000460680.6", "NM_004656.4", "GRCh38", "Waters et al. 2024, Nat. Genet. 56:1434–1445", "CC BY 4.0"],
]
S1 = w("Table_S1.tsv", s1_header, s1_rows)

# ---- Table S2: predictor coverage by region (full_model_coverage.tsv) -------
cov = pd.read_csv(T / "full_model_coverage.tsv", sep="\t")
cov_regions = ["all", "splice", "intronic", "coding", "missense"]
s2_header = ["predictor"] + [f"{r} (scored/total)" for r in cov_regions]
s2_rows = []
for m in ORDER:
    row = [NAME[m]]
    for r in cov_regions:
        c = cov[(cov.model == m) & (cov.region == r)]
        if len(c):
            row.append(f"{int(c.n_scored.iloc[0])}/{int(c.n_total.iloc[0])}")
        else:
            row.append("—")
    s2_rows.append(row)
S2 = w("Table_S2.tsv", s2_header, s2_rows)

# ---- Table S3: a priori power / precision (precision_estimate.tsv, POOLED) ---
prec = pd.read_csv(T / "precision_estimate.tsv", sep="\t")
pooled = prec[prec.gene == "POOLED"]
s3_header = ["region_subset", "N_spearman", "spearman_halfwidth_rho0.5",
             "n_path", "n_benign", "auroc_halfwidth_0.85",
             "verdict_spearman", "verdict_auroc"]
s3_rows = []
for _, r in pooled.iterrows():
    s3_rows.append([r.region_subset, int(r.N_spearman), r["sp_hw_rho0.5"],
                    int(r.n_pos), int(r.n_neg),
                    r["auroc_hw_0.85"] if str(r["auroc_hw_0.85"]) not in ("nan", "") else "—",
                    r.verdict_spearman, r.verdict_auroc])
S3 = w("Table_S3.tsv", s3_header, s3_rows)

# ---- Table S4: per-gene Spearman rho on splice (spearman_by_gene.tsv) --------
bg = pd.read_csv(T / "spearman_by_gene.tsv", sep="\t")
spg = bg[bg.region == "splice"]
s4_header = ["predictor"] + GENES
s4_rows = []
for m in ORDER:
    row = [NAME[m]]
    for g in GENES:
        c = spg[(spg.model == m) & (spg.gene == g)]
        val = c.rho.iloc[0] if len(c) else ""
        row.append(val if str(val) not in ("", "nan") else "—")
    s4_rows.append(row)
S4 = w("Table_S4.tsv", s4_header, s4_rows)

# ---- Table S5: ClinVar binary AUROC & AUPRC on splice (auroc_clinvar.tsv) ----
au = pd.read_csv(T / "auroc_clinvar.tsv", sep="\t")


def cell(model, region, version, metric):
    r = au[(au.model == model) & (au.region == region) & (au.version == version)]
    if not len(r) or str(r[metric].iloc[0]) in ("", "nan"):
        return "—"
    v = r.iloc[0]
    lo, hi = v[f"{metric}_lo"], v[f"{metric}_hi"]
    return f"{v[metric]:.3f} ({lo:.3f}–{hi:.3f})"


s5_header = ["predictor", "AUROC splice (incl BRCA1)", "AUROC splice (excl BRCA1)",
             "AUPRC splice (incl BRCA1)", "AUPRC splice (excl BRCA1)"]
s5_rows = []
for m in ORDER:
    s5_rows.append([NAME[m],
                    cell(m, "splice", "incl_BRCA1", "auroc"),
                    cell(m, "splice", "excl_BRCA1", "auroc"),
                    cell(m, "splice", "incl_BRCA1", "auprc"),
                    cell(m, "splice", "excl_BRCA1", "auprc")])
S5 = w("Table_S5.tsv", s5_header, s5_rows)

# =============================== DOCX =======================================
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(10)

title = doc.add_heading("Supplementary Materials", level=0)
doc.add_paragraph(
    "Functional-evidence benchmarking of splice- and sequence-based variant effect "
    "predictors in hereditary cancer genes. All values are reproduced from the "
    "released results tables (results/tables/) and the project data record; nothing "
    "here is recomputed or estimated independently of the main analysis."
)

CAPTIONS = {
 "S1": ("Supplementary Table S1. Functional gold-standard (SGE) score sets used per "
        "gene (referenced in Methods 2.2 and Data availability). All seven are "
        "saturation genome editing (SGE) and openly licensed (CC0 or CC BY 4.0). "
        "ENST-based sets were mapped to GRCh38 via the listed MANE-Select RefSeq "
        "transcript (validated 12/12 per gene). Primary publications are author-verified "
        "(PubMed) and cross-checked with the manuscript reference list where applicable; "
        "BARD1 and PALB2 have no peer-reviewed primary publication (BARD1: medRxiv "
        "preprint only), so the MaveDB URN is the authoritative source for those two."),
 "S2": ("Supplementary Table S2. Predictor coverage by region (referenced in Methods "
        "2.3). Cells are variants scored / total in each region. AlphaMissense scores "
        "1/1,781 splice variants (protein/missense only); gnomAD covers only variants "
        "recorded in gnomAD v4; all DNA/splice models and conservation cover every "
        "variant. Evo2/ESM were not scored and are omitted."),
 "S3": ("Supplementary Table S3. A priori power / precision analysis, pooled "
        "(referenced in §4 Limitations). Computed in Milestone 3 on the then-5-gene "
        "panel (BRCA1, BRCA2, BARD1, PALB2, RAD51C) over full functional-SNV "
        "denominators; the final 7-gene panel only increases power. Half-widths are "
        "95% CI half-widths at the stated assumed effect size; 'sufficient' = "
        "Spearman half-width ≤0.10 / AUROC ≤0.05. The pooled splice and whole-gap "
        "continuous analyses are sufficiently powered."),
 "S4": ("Supplementary Table S4. Per-gene Spearman ρ of each predictor vs the "
        "continuous functional gold standard on the splice subset (the per-gene "
        "values pooled in main Table 2 and shown in Fig. 2). '—' = not estimable "
        "(e.g. AlphaMissense, n=1 on splice; or <4 scored variants in a gene)."),
 "S5": ("Supplementary Table S5. ClinVar binary AUROC and AUPRC on the splice subset, "
        "with and without BRCA1 (the circularity sensitivity in Results 3.4), as "
        "value (95% CI). Removing BRCA1 (the only circularity-flagged gene) leaves "
        "the splice metrics essentially unchanged; every splice-aware model saturates "
        "near 1.0, which is why the continuous functional correlation, not this "
        "saturated binary metric, is the primary result. AlphaMissense is omitted on "
        "splice (n too few)."),
}


def add_table(key, header, rows, fontsz=8):
    cap = doc.add_paragraph()
    r = cap.add_run(CAPTIONS[key].split(".")[0] + ".")
    r.bold = True
    cap.add_run(" " + CAPTIONS[key].split(".", 1)[1].strip())
    for run in cap.runs:
        run.font.size = Pt(9)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        rr = c.paragraphs[0].add_run(str(h))
        rr.bold = True
        rr.font.size = Pt(fontsz)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            rr = cells[j].paragraphs[0].add_run(str(v))
            rr.font.size = Pt(fontsz)
    doc.add_paragraph()


add_table("S1", *S1, fontsz=8)
add_table("S2", *S2)
add_table("S3", *S3)
add_table("S4", *S4)
add_table("S5", *S5, fontsz=7)

doc.save(str(OUTDIR / "Supplementary_Materials.docx"))
print("Wrote to", OUTDIR)
for n in ["Table_S1.tsv", "Table_S2.tsv", "Table_S3.tsv", "Table_S4.tsv",
          "Table_S5.tsv", "Supplementary_Materials.docx"]:
    print("  -", n)
