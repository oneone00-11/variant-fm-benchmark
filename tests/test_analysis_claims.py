"""Every analysis the manuscript claims must have code and output behind it.

A value check asks whether a number is right. It cannot ask whether the analysis
a sentence describes exists at all. The sampling-frame paragraph reported a
whole-gene enrichment of 0.713 versus 0.578; both numbers were plausible, both
were near real values elsewhere in the pipeline, and the comparison had never
been computed -- ipw_frame_balance.csv is restricted to the splice window.

Claims bind by content anchor, never by paragraph number: rewriting the abstract
shifted every later paragraph by two and left all six `para` values pointing at
the wrong text, with nothing to notice.
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "phase1"))

MANIFEST = REPO / "phase1" / "config" / "analysis_claims.json"
def _manuscript() -> Path | None:
    """The .docx these claims bind to, if it is on this machine.

    Resolved from $VARIANT_FM_MANUSCRIPT, else the newest calibration_draft*.docx
    (or legacy draft_reframed_*.docx) on the Desktop. The filename is not
    hard-coded: it carries a working title that is nobody's business but the
    author's, and the tests only need some draft to check the anchors against.
    """
    env = os.environ.get("VARIANT_FM_MANUSCRIPT")
    if env:
        p = Path(env).expanduser()
        return p if p.exists() else None
    desk = Path.home() / "Desktop"
    hits = sorted([*desk.glob("calibration_draft*.docx"),
                   *(desk / "calibration").glob("calibration_draft*.docx"),
                   *desk.glob("draft_reframed_*.docx")],
                  key=lambda q: q.stat().st_mtime, reverse=True)
    return hits[0] if hits else None


MANUSCRIPT = _manuscript()
VALID_STATUS = {"verified", "narrative", "gap"}


def _claims() -> list[dict]:
    assert MANIFEST.exists(), f"{MANIFEST.relative_to(REPO)} is missing"
    return json.loads(MANIFEST.read_text())["claims"]


def _outputs(c) -> list[str]:
    out = c.get("output")
    return [] if not out else (out if isinstance(out, list) else [out])


def test_manifest_is_well_formed():
    claims = _claims()
    assert claims, "no claims recorded"
    for c in claims:
        assert c.get("claim", "").strip(), c
        assert c.get("status") in VALID_STATUS, c
        assert c.get("anchor"), f"no content anchor: {c['claim']}"


def test_every_named_script_and_output_exists():
    missing = [f"{c['claim']}: {rel}" for c in _claims()
               for rel in ([c["script"]] if c.get("script") else []) + _outputs(c)
               if not (REPO / rel).exists()]
    assert not missing, "claims naming files that do not exist:\n  " + "\n  ".join(missing)


def test_every_anchor_hits_exactly_once():
    """Zero hits means the sentence was reworded and the row now describes
    nothing; more than one means the anchor could bind to the wrong sentence.
    Both are invisible under paragraph numbers, which always resolve to *some*
    paragraph however far the text has moved."""
    if MANUSCRIPT is None:
        pytest.skip("manuscript not available in this checkout")
    from src.check_manuscript_numbers import _blocks

    blocks = _blocks(MANUSCRIPT)
    bad = [f"{c['para']}: {sum(t.count(c['anchor']) for _, t in blocks)} hits "
           f"for {c['anchor']!r}" for c in _claims()
           if sum(t.count(c["anchor"]) for _, t in blocks) != 1]
    assert not bad, "anchors that do not hit exactly once:\n  " + "\n  ".join(bad)


def test_paragraph_numbers_are_not_load_bearing(tmp_path):
    """Regression for the abstract rewrite: with every `para` deliberately
    wrong, resolution must be unaffected, because nothing matches on it."""
    if MANUSCRIPT is None:
        pytest.skip("manuscript not available in this checkout")
    import src.check_manuscript_numbers as cmn

    before = cmn.resolve_anchors(MANUSCRIPT)
    spec = json.loads(MANIFEST.read_text())
    for c in spec["claims"]:
        c["para"] = f"P{int(c['para'][1:]) + 23}"
    shifted = tmp_path / "shifted.json"
    shifted.write_text(json.dumps(spec))
    assert before == cmn.resolve_anchors(MANUSCRIPT, shifted) and before


# --- the scoped rule, after the global-fallback escape hatch was removed ------

PRE_FIX_FIXTURE = Path(__file__).parent / "fixtures" / "pre_fix_clinical_yield.docx"


def _backup_with_the_defect():
    """The clinical-yield paragraph as it stood before 74.5 was corrected.

    Prefers the committed fixture, which carries that paragraph verbatim, so the
    check runs for everyone. This used to glob the author's Desktop for the
    pre-fix draft and skip when it was absent -- which meant the suite reported
    63 passed on that one machine and 62 passed 1 skipped everywhere else, and
    the disclosed count was true of nobody but the author. The full draft is
    still used when it happens to be present, as a check on the fixture.
    """
    if PRE_FIX_FIXTURE.exists():
        return PRE_FIX_FIXTURE
    hits = sorted((Path.home() / "Desktop").glob(
        "**/draft_reframed_*_backup_*preMerge*.docx"))
    return hits[-1] if hits else None


def test_a_number_absent_from_its_own_declared_source_is_reported(tmp_path, monkeypatch):
    """The regression this rule exists for.

    74.5 was absent from the yield table its paragraph declares, matched exactly
    one unrelated value in the global pool, and passed under the old rule while
    being the wrong number for its own sentence. Synthetic so it runs anywhere.
    """
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_paragraph("padding")
    d.add_paragraph("The share moved from 17.7% to 74.5% under the new definition.")
    doc = tmp_path / "m.docx"
    d.save(doc)

    results = tmp_path / "results"
    results.mkdir()
    # declared source holds 17.7 but not 74.5; a second file holds 74.5 once, so
    # the old rule would have called it "a specific source, just not this one"
    (results / "declared.csv").write_text("a\n0.177\n")
    (results / "unrelated.csv").write_text("b\n74.5\n")
    claims = tmp_path / "claims.json"
    claims.write_text(json.dumps({"claims": [
        {"para": "P1", "claim": "share under the new definition",
         "anchor": "The share moved from", "script": None,
         "output": "results/declared.csv", "status": "verified"}]}))

    # scoped paths resolve against the repository root, so point it here
    monkeypatch.setattr(cmn, "REPO", tmp_path)
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims)
    flagged = {t["token"] for t in r["scoped_mismatch"]}
    assert "74.5" in flagged, r["scoped_mismatch"]
    assert "17.7" not in flagged, "a value present in the declared source must pass"


def test_the_rule_reports_it_on_the_actual_pre_fix_text():
    """The synthetic case above proves the rule; this proves it on the real text.

    Runs against the committed fixture, so it does not depend on any file
    outside the repository.
    """
    backup = _backup_with_the_defect()
    assert backup is not None, (
        f"neither the fixture {PRE_FIX_FIXTURE.name} nor a pre-fix draft is "
        "available; the fixture is tracked and should always be present")
    import src.check_manuscript_numbers as cmn

    r = cmn.classify(backup)
    tokens = {t["token"] for t in r["scoped_mismatch"]}
    assert "74.5" in tokens, sorted(tokens)
    assert "17.7" not in tokens, "a value present in the declared source must pass"


def test_a_declared_derivation_licenses_exactly_one_value(tmp_path, monkeypatch):
    """Route (b): a summary of stored values carries its arithmetic. It must
    admit that value and nothing else, so it cannot become a paragraph pass."""
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_paragraph("padding")
    d.add_paragraph("There were 30 disagreements and 41 agreements.")
    doc = tmp_path / "m.docx"
    d.save(doc)
    results = tmp_path / "results"
    results.mkdir()
    (results / "declared.csv").write_text("a\n18\n12\n")
    claims = tmp_path / "claims.json"
    claims.write_text(json.dumps({"claims": [
        {"para": "P1", "claim": "disagreements", "anchor": "There were 30 disagreements",
         "script": None, "output": "results/declared.csv", "status": "verified",
         "derived": [{"value": 30, "how": "18 + 12"}]}]}))

    monkeypatch.setattr(cmn, "REPO", tmp_path)
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims)
    flagged = {t["token"] for t in r["scoped_mismatch"]}
    assert "30" not in flagged          # licensed by its declared arithmetic
    assert "41" in flagged              # not licensed, still reported


def test_the_section_number_exemption_covers_headings_not_measurements(tmp_path, monkeypatch):
    """The whitelist exempts section numbers by pattern (a digit or two, a point, a
    digit), and a pattern cannot tell the cross-reference "2.4" from the measurement
    "8.8". The v1 -> v2 replacement pass skipped a stale Table 3 value, 8.8%, as a
    section number. The exemption now holds only for numbers that open a heading."""
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_heading("2.4 Metrics and inference", level=2)
    d.add_paragraph("As set out in 2.4, the high-confidence share was 8.8% of variants.")
    doc = tmp_path / "m.docx"
    d.save(doc)
    results = tmp_path / "results"
    results.mkdir()
    (results / "t.csv").write_text("a\n0.5\n")
    wl = tmp_path / "wl.json"
    wl.write_text(json.dumps({"section_numbers": {"reason": "headings",
                                                  "pattern": r"^(?:[1-9]|1[0-9])\.[0-9]$"}}))
    monkeypatch.setattr(cmn, "REPO", tmp_path)
    r = cmn.classify(doc, results=results, whitelist=wl, claims=tmp_path / "none.json")
    # the cross-reference is set aside -- by the kind rules as a section reference, or
    # by the whitelist pattern when no kinds file applies
    exempt = {t["token"] for t in r["whitelisted"] + r.get("non_measurement", [])}
    unsourced = {t["token"] for t in r["no_source"]}
    assert "2.4" in exempt, "a cross-reference to a real heading stays exempt"
    assert "8.8" in unsourced and "8.8" not in exempt, "a measurement shaped like a section number is not"



# --- binding to named cells, and the rules that set non-measurements aside -----

def _kinds_file(tmp_path, rules):
    k = tmp_path / "kinds.json"
    k.write_text(json.dumps({"rules": rules}))
    return k


LABEL_RULE = {"name": "figure and table labels", "kind": "figure or table label",
              "pattern": r"\b(?:Figures?|Tables?)\s+(\d+)", "reason": "labels"}


def test_a_caption_claim_binds_its_table_to_named_cells(tmp_path, monkeypatch):
    """A claim anchored on a table caption with "table": true scopes the table below it.
    A scoped number is BOUND only when a named cell of a declared output holds it, and
    the binding records that cell; an unscoped number is only a pool match."""
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_paragraph("padding")
    d.add_paragraph("Table 9. Brier by object.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Object", "Brier"
    t.cell(1, 0).text, t.cell(1, 1).text = "fusion", "0.063"
    d.add_paragraph("After the table the fusion's Brier was 0.063 and 0.071.")
    doc = tmp_path / "m.docx"
    d.save(doc)
    results = tmp_path / "results"
    results.mkdir()
    (results / "summary.csv").write_text("model,Brier\nfusion,0.0630194\nsingle,0.0733\n")
    claims = tmp_path / "claims.json"
    claims.write_text(json.dumps({"claims": [
        {"para": "P1", "claim": "Table 9", "anchor": "Table 9. Brier by object", "script": None,
         "output": "results/summary.csv", "status": "verified", "table": True}]}))
    monkeypatch.setattr(cmn, "REPO", tmp_path)
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims,
                     kinds=_kinds_file(tmp_path, [LABEL_RULE]))
    cell = [x for x in r["bound"] if x["uid"] == "T1r1c1"]
    assert cell and cell[0]["cells"] == ["summary.csv[fusion].Brier = 0.0630194"]
    assert [(x["uid"], x["token"]) for x in r["pool"]] == [("P2", "0.063")]
    assert [x["token"] for x in r["no_source"]] == ["0.071"]
    assert [x["token"] for x in r["non_measurement"]] == ["9"]
    assert r["coverage"] == {"tokens": 4, "measured": 3, "bound": 1, "single_cell": 1, "several_cells": 0}


def test_a_parameter_rule_whose_code_no_longer_defines_the_value_fails(tmp_path, monkeypatch):
    """A parameter is set aside only while the code still sets it: the rule names a file
    and a pattern that must match there."""
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_paragraph("padding")
    d.add_paragraph("Intervals use 2,000 resamples.")
    doc = tmp_path / "m.docx"
    d.save(doc)
    results = tmp_path / "results"
    results.mkdir()
    (results / "t.csv").write_text("a\n0.5\n")
    (tmp_path / "code.py").write_text("N_BOOT = 1000\n")
    rule = {"name": "bootstrap resamples", "kind": "parameter", "pattern": r"\b(2,000) resamples",
            "reason": "resample count", "source": "code.py", "defined_by": r"(?m)^N_BOOT = 2000\b"}
    monkeypatch.setattr(cmn, "REPO", tmp_path)
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=tmp_path / "none.json",
                     kinds=_kinds_file(tmp_path, [rule]))
    assert [q["rule"] for q in r["rule_problems"]] == ["bootstrap resamples"]
    assert [t["token"] for t in r["non_measurement"]] == ["2,000"]


def test_every_parameter_rule_names_a_live_definition():
    """The manuscript's parameters and constants are set aside because code sets them;
    every such rule must still resolve in the repository."""
    import re
    import src.check_manuscript_numbers as cmn

    rules = json.loads(cmn.KINDS.read_text())["rules"]
    assert rules
    for r in rules:
        re.compile(r["pattern"])
        if r["kind"] in ("parameter", "literature constant", "derived constant"):
            assert r.get("source") and r.get("defined_by"), f"{r['name']}: a {r['kind']} rule must name its source"
            assert re.search(r["defined_by"], (REPO / r["source"]).read_text()), \
                f"{r['name']}: {r['source']} no longer defines the value"


def test_a_computed_derivation_binds_only_the_value_it_computes(tmp_path, monkeypatch):
    """A derivation with `agg` is evaluated against the output it names; the token binds
    only if the computed value is the printed one. A count that drifts is caught."""
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_paragraph("padding")
    d.add_paragraph("The interval excluded zero in 2 comparisons; by weight phyloP ranks 3.")
    doc = tmp_path / "m.docx"
    d.save(doc)
    results = tmp_path / "results"
    results.mkdir()
    table = results / "t.csv"
    table.write_text("tool,excl,w\nfusion,True,0.9\nphylop,False,0.1\ncadd,True,0.5\n")
    claims = tmp_path / "claims.json"
    claims.write_text(json.dumps({"claims": [
        {"para": "P1", "claim": "counts and ranks", "anchor": "The interval excluded zero in", "script": None,
         "output": "results/t.csv", "status": "verified",
         "derived": [{"value": 2, "agg": "count", "from": "results/t.csv", "where": {"excl": "True"}, "how": "count"},
                     {"value": 3, "agg": "rank", "from": "results/t.csv", "column": "w", "key": "phylop", "how": "rank"}]}]}))
    monkeypatch.setattr(cmn, "REPO", tmp_path)
    kinds = tmp_path / "kinds.json"
    kinds.write_text(json.dumps({"rules": []}))
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims, kinds=kinds)
    got = {t["token"]: t["cells"][0] for t in r["bound"]}
    assert got["2"] == "derived: count = 2" and got["3"] == "derived: rank = 3"

    table.write_text("tool,excl,w\nfusion,True,0.9\nphylop,False,0.1\ncadd,False,0.5\n")
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims, kinds=kinds)
    assert "2" in {t["token"] for t in r["scoped_mismatch"]}, "the count is now 1; the printed 2 must fail"


def test_a_cell_binds_only_if_it_rounds_to_the_printed_number(tmp_path, monkeypatch):
    """A window of half a unit either side of 0.997 admits BRCA2's control AUROC 0.9975, which
    prints as 0.998 -- the same coincidence, from the other side, that once let a printed 0.998
    pass for TP53's 0.9969. A cell binds only if it rounds to the printed number."""
    import docx as docxlib
    import src.check_manuscript_numbers as cmn

    d = docxlib.Document()
    d.add_paragraph("padding")
    d.add_paragraph("Post-orientation control AUROC 0.997 in TP53.")
    doc = tmp_path / "m.docx"
    d.save(doc)
    results = tmp_path / "results"
    results.mkdir()
    table = results / "gate.csv"
    claims = tmp_path / "claims.json"
    claims.write_text(json.dumps({"claims": [
        {"para": "P1", "claim": "orientation gate", "anchor": "Post-orientation control AUROC", "script": None,
         "output": "results/gate.csv", "status": "verified"}]}))
    monkeypatch.setattr(cmn, "REPO", tmp_path)
    kinds = _kinds_file(tmp_path, [])
    table.write_text("gene,control_auroc\nBRCA2,0.9975\n")
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims, kinds=kinds)
    assert [t["token"] for t in r["scoped_mismatch"]] == ["0.997"], "0.9975 prints as 0.998; it must not bind 0.997"
    table.write_text("gene,control_auroc\nBRCA2,0.9975\nTP53,0.996925\n")
    r = cmn.classify(doc, results=results, whitelist=tmp_path / "none.json", claims=claims, kinds=kinds)
    assert [t["cells"] for t in r["bound"]] == [["gate.csv[TP53].control_auroc = 0.996925"]]
